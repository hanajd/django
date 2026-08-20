#!/usr/bin/env python3
"""Convert Word (.docx) documents to Markdown."""

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path

from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

try:
    from docx import Document
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "python-docx is required. Install with: pip install python-docx"
    ) from exc

HEADING_STYLE_LEVELS = {
    "Title": 1,
    "Heading 1": 1,
    "Heading 2": 2,
    "Heading 3": 3,
    "Heading 4": 4,
    "Heading 5": 5,
    "Heading 6": 6,
    "Caption": 4,
}

# Word form checkboxes often use Wingdings 2 via w:sym (not exposed in run.text).
WINGDINGS2_CHECKBOX_MAP = {
    ("Wingdings 2", "00A3"): "□",
    ("Wingdings 2", "0052"): "☑",
}


def _sym_to_char(font: str, char_hex: str) -> str:
    key = (font.strip(), char_hex.upper())
    if key in WINGDINGS2_CHECKBOX_MAP:
        return WINGDINGS2_CHECKBOX_MAP[key]
    return ""


def _run_inline_text(run) -> str:
    parts: list[str] = []
    for child in run._element:
        if child.tag == qn("w:t"):
            parts.append(child.text or "")
        elif child.tag == qn("w:sym"):
            font = child.get(qn("w:font"), "")
            char = child.get(qn("w:char"), "")
            sym = _sym_to_char(font, char)
            if sym:
                parts.append(sym)
        elif child.tag == qn("w:tab"):
            parts.append("\t")
    return "".join(parts)


def _paragraph_plain_text(paragraph: Paragraph) -> str:
    return "".join(_run_inline_text(run) for run in paragraph.runs)


@dataclass
class TableAnchor:
    row: int
    col: int
    colspan: int
    rowspan: int
    text: str


def iter_block_items(document: DocumentObject):
    """Yield paragraphs and tables in document order."""
    parent = document.element.body
    for child in parent.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _unique_row_cells(row) -> list[_Cell]:
    seen: set[int] = set()
    cells: list[_Cell] = []
    for cell in row.cells:
        tc_id = id(cell._tc)
        if tc_id in seen:
            continue
        seen.add(tc_id)
        cells.append(cell)
    return cells


def _cell_colspan(cell: _Cell) -> int:
    tc_pr = cell._tc.find(qn("w:tcPr"))
    if tc_pr is None:
        return 1
    grid_span = tc_pr.find(qn("w:gridSpan"))
    if grid_span is None:
        return 1
    return max(1, int(grid_span.get(qn("w:val"), 1)))


def _cell_text(cell: _Cell) -> str:
    text = "\n".join(_paragraph_plain_text(p) for p in cell.paragraphs)
    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    text = re.sub(r"\n+", " ", text)
    return text


def _cell_image_refs(cell: _Cell, image_map: dict[str, str]) -> list[str]:
    refs: list[str] = []
    seen: set[str] = set()
    for paragraph in cell.paragraphs:
        for blip in paragraph._element.findall(".//" + qn("a:blip")):
            embed = blip.get(qn("r:embed"))
            if not embed or embed in seen:
                continue
            seen.add(embed)
            path = image_map.get(embed)
            if path:
                refs.append(f"![]({path})")
    return refs


def _cell_content(cell: _Cell, image_map: dict[str, str]) -> str:
    text = _cell_text(cell)
    images = _cell_image_refs(cell, image_map)
    parts: list[str] = []
    if text:
        parts.append(text)
    parts.extend(images)
    if not parts:
        return ""
    if len(parts) == 1:
        return parts[0]
    return "\n".join(parts)


def table_has_cell_images(table: Table, image_map: dict[str, str]) -> bool:
    for row in table.rows:
        for cell in _unique_row_cells(row):
            if _cell_image_refs(cell, image_map):
                return True
    return False


def _escape_md_table_cell(text: str) -> str:
    text = text.replace("|", r"\|")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def table_has_merges(table: Table) -> bool:
    seen_tc_by_row: list[set[int]] = []
    tc_first_row: dict[int, int] = {}

    for ri, row in enumerate(table.rows):
        seen: set[int] = set()
        for cell in row.cells:
            tc_id = id(cell._tc)
            if tc_id in seen:
                return True
            seen.add(tc_id)
            if _cell_colspan(cell) > 1:
                return True
            if tc_id in tc_first_row and tc_first_row[tc_id] != ri:
                return True
            tc_first_row.setdefault(tc_id, ri)
        seen_tc_by_row.append(seen)
    return False


def _iter_row_cell_spans(row) -> list[tuple[int, _Cell, int]]:
    spans: list[tuple[int, _Cell, int]] = []
    col = 0
    seen: set[int] = set()
    for cell in row.cells:
        tc_id = id(cell._tc)
        if tc_id in seen:
            continue
        seen.add(tc_id)
        colspan = _cell_colspan(cell)
        spans.append((col, cell, colspan))
        col += colspan
    return spans


def _cell_at_grid_column(table: Table, row_idx: int, col_idx: int) -> _Cell | None:
    for col, cell, colspan in _iter_row_cell_spans(table.rows[row_idx]):
        if col <= col_idx < col + colspan:
            return cell
    return None


def _rowspan_from(table: Table, row_idx: int, col_idx: int, cell: _Cell) -> int:
    span = 1
    ref = id(cell._tc)
    for ri in range(row_idx + 1, len(table.rows)):
        below = _cell_at_grid_column(table, ri, col_idx)
        if below is None or id(below._tc) != ref:
            break
        span += 1
    return span


def _analyze_table_anchors(
    table: Table,
    image_map: dict[str, str] | None = None,
) -> tuple[int, int, dict[tuple[int, int], TableAnchor]]:
    nrows = len(table.rows)
    ncols = len(table.columns)
    grid_tc: list[list[int | None]] = [[None] * ncols for _ in range(nrows)]
    anchors: dict[tuple[int, int], TableAnchor] = {}

    for ri, row in enumerate(table.rows):
        for ci, cell, colspan in _iter_row_cell_spans(row):
            if ci >= ncols or grid_tc[ri][ci] is not None:
                continue

            colspan = min(colspan, ncols - ci)
            rowspan = min(_rowspan_from(table, ri, ci, cell), nrows - ri)
            tc_id = id(cell._tc)

            anchor = TableAnchor(
                row=ri,
                col=ci,
                colspan=colspan,
                rowspan=rowspan,
                text=_cell_content(cell, image_map) if image_map else _cell_text(cell),
            )
            anchors[(ri, ci)] = anchor
            for r in range(ri, ri + rowspan):
                for c in range(ci, ci + colspan):
                    if r < nrows and c < ncols:
                        grid_tc[r][c] = tc_id

    return nrows, ncols, anchors


def _escape_html(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def _cell_html_content(content: str) -> str:
    parts: list[str] = []
    for line in content.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith("!["):
            parts.append(line)
        else:
            parts.append(_escape_html(line))
    return "<br/>".join(parts)


def table_to_html(table: Table, image_map: dict[str, str] | None = None) -> str:
    nrows, ncols, anchors = _analyze_table_anchors(table, image_map)
    if not anchors:
        return ""

    skip: set[tuple[int, int]] = set()
    for anchor in anchors.values():
        if anchor.rowspan > 1:
            for r in range(anchor.row + 1, anchor.row + anchor.rowspan):
                for c in range(anchor.col, anchor.col + anchor.colspan):
                    skip.add((r, c))

    lines = ["<table>"]
    for ri in range(nrows):
        cells: list[str] = []
        ci = 0
        while ci < ncols:
            if (ri, ci) in skip:
                ci += 1
                continue
            anchor = anchors.get((ri, ci))
            if anchor is None:
                ci += 1
                continue
            attrs: list[str] = []
            if anchor.colspan > 1:
                attrs.append(f'colspan="{anchor.colspan}"')
            if anchor.rowspan > 1:
                attrs.append(f'rowspan="{anchor.rowspan}"')
            attr_str = f" {' '.join(attrs)}" if attrs else ""
            cells.append(f"<td{attr_str}>{_cell_html_content(anchor.text)}</td>")
            ci += anchor.colspan
        if cells:
            lines.append(f"<tr>{''.join(cells)}</tr>")
    lines.append("</table>")
    return "\n".join(lines)


def _runs_to_markdown(paragraph: Paragraph) -> str:
    parts: list[str] = []
    for run in paragraph.runs:
        text = _run_inline_text(run)
        if not text:
            continue
        if run.bold and run.italic:
            text = f"***{text}***"
        elif run.bold:
            text = f"**{text}**"
        elif run.italic:
            text = f"*{text}*"
        parts.append(text)
    return "".join(parts).strip()


def _paragraph_image_refs(paragraph: Paragraph, image_map: dict[str, str]) -> list[str]:
    refs: list[str] = []
    seen: set[str] = set()
    for blip in paragraph._element.findall(".//" + qn("a:blip")):
        embed = blip.get(qn("r:embed"))
        if not embed or embed in seen:
            continue
        seen.add(embed)
        path = image_map.get(embed)
        if path:
            refs.append(f"![]({path})")
    return refs


def _plain_paragraph_text(text: str) -> str:
    return re.sub(r"\*+", "", text).strip()


def _is_figure_caption_text(text: str) -> bool:
    plain = _plain_paragraph_text(text)
    return bool(plain) and plain.endswith("图") and len(plain) >= 3


def paragraph_to_markdown(
    paragraph: Paragraph,
    image_map: dict[str, str],
    *,
    heading_before_table: bool = False,
) -> str:
    style_name = paragraph.style.name if paragraph.style else "Normal"
    text = _runs_to_markdown(paragraph)
    image_lines = _paragraph_image_refs(paragraph, image_map)

    if not text and not image_lines:
        return ""

    lines: list[str] = []
    caption_after_images = bool(text and image_lines and _is_figure_caption_text(text))

    if text and not caption_after_images:
        level = HEADING_STYLE_LEVELS.get(style_name)
        if level and not heading_before_table:
            lines.append(f"{'#' * level} {text}")
        else:
            lines.append(text)

    lines.extend(image_lines)

    if text and caption_after_images:
        lines.append(text)

    return "\n\n".join(lines)


def table_to_markdown(
    table: Table,
    image_map: dict[str, str],
    *,
    force_simple: bool = False,
) -> str:
    has_images = table_has_cell_images(table, image_map)
    if not force_simple and (table_has_merges(table) or has_images):
        return table_to_html(table, image_map)

    rows: list[list[str]] = []
    for row in table.rows:
        cells = [
            _escape_md_table_cell(_cell_content(cell, image_map))
            for cell in _unique_row_cells(row)
        ]
        if any(cells):
            rows.append(cells)

    if not rows:
        return ""

    max_cols = max(len(row) for row in rows)
    normalized = [row + [""] * (max_cols - len(row)) for row in rows]
    header = normalized[0]
    body = normalized[1:] if len(normalized) > 1 else []

    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * max_cols) + " |",
    ]
    for row in body:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def build_image_map(
    document: DocumentObject,
    images_dir: Path,
    markdown_dir: Path,
) -> dict[str, str]:
    images_dir.mkdir(parents=True, exist_ok=True)
    image_map: dict[str, str] = {}
    used_names: dict[str, int] = {}

    for rel_id, rel in document.part.rels.items():
        if "image" not in rel.reltype:
            continue
        part = rel.target_part
        blob = part.blob
        ext = Path(part.partname).suffix or ".png"
        base = Path(part.partname).name
        if base in used_names:
            used_names[base] += 1
            base = f"{Path(base).stem}_{used_names[base]}{ext}"
        else:
            used_names[base] = 1
        out_path = images_dir / base
        out_path.write_bytes(blob)
        rel_path = out_path.relative_to(markdown_dir).as_posix()
        image_map[rel_id] = rel_path
    return image_map


def convert_docx_to_markdown(
    docx_path: Path,
    output_path: Path | None = None,
    images_dir: Path | None = None,
    simple_tables: bool = False,
) -> str:
    docx_path = docx_path.resolve()
    if output_path is None:
        output_path = docx_path.with_suffix(".md")
    else:
        output_path = output_path.resolve()

    if images_dir is None:
        images_dir = output_path.parent / "images"
    else:
        images_dir = images_dir.resolve()

    document = Document(str(docx_path))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    image_map = build_image_map(document, images_dir, output_path.parent)

    blocks_raw = list(iter_block_items(document))
    blocks: list[str] = []
    for idx, block in enumerate(blocks_raw):
        if isinstance(block, Paragraph):
            md = paragraph_to_markdown(
                block,
                image_map,
                heading_before_table=(
                    idx + 1 < len(blocks_raw) and isinstance(blocks_raw[idx + 1], Table)
                ),
            )
            if md:
                blocks.append(md)
        elif isinstance(block, Table):
            md = table_to_markdown(block, image_map, force_simple=simple_tables)
            if md:
                blocks.append(md)

    markdown = "\n\n".join(blocks).strip() + "\n"
    output_path.write_text(markdown, encoding="utf-8")
    return markdown


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert Word (.docx) to Markdown.")
    parser.add_argument("input", help="Input .docx file")
    parser.add_argument(
        "-o",
        "--output",
        help="Output .md file (default: same name as input)",
    )
    parser.add_argument(
        "--images-dir",
        help="Directory for extracted images (default: <output-dir>/images)",
    )
    parser.add_argument(
        "--simple-tables",
        action="store_true",
        help="Always use Markdown pipe tables (merged cells become HTML by default)",
    )
    args = parser.parse_args()

    input_path = Path(args.input).resolve()
    if not input_path.is_file():
        parser.error(f"Input file not found: {input_path}")
    if input_path.suffix.lower() != ".docx":
        parser.error("Input must be a .docx file")

    output_path = Path(args.output).resolve() if args.output else None
    images_dir = Path(args.images_dir).resolve() if args.images_dir else None

    convert_docx_to_markdown(
        input_path,
        output_path,
        images_dir,
        simple_tables=args.simple_tables,
    )
    out = output_path or input_path.with_suffix(".md")
    print(f"Written to {out}")
    img_dir = images_dir or out.parent / "images"
    if img_dir.is_dir():
        count = len(list(img_dir.glob("*")))
        print(f"Extracted {count} image(s) to {img_dir}")


if __name__ == "__main__":
    main()
